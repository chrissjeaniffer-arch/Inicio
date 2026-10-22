import sys, json, datetime as dt
sys.path.insert(0, sys.argv[3])
from content import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

plan = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
items, links, names = plan["items"], plan["links"], plan["names"]

AZUL, SALMON, GRIS = "004899", "E6375D", "4B4B4D"
F = "Montserrat"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=AZUL)
body = Font(name=F, size=9, color=GRIS)
bold = Font(name=F, size=9, color=AZUL, bold=True)
thin = Side(style="thin", color="D0D5DD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)
band = PatternFill("solid", fgColor="EEF1F5")
LEVEL_FILL = {"Crítica": "FBD5D6", "Alta": "FDEBC8", "Media": "D6ECFE", "Baja": "D8F0EE"}
STATE_FILL = {"Resuelto en la propuesta": "D8F0EE", "Condicionado a validación": "FDEBC8",
              "Pendiente de decisión": "D6ECFE", "No resuelto": "FBD5D6"}

wb = Workbook()


def D(s):
    return dt.date.fromisoformat(s) if s else None


def sheet(title, headers, rows, widths, date_cols=(), first_bold=True, note=None, tab=AZUL):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=F, size=10, italic=True, color=GRIS)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 30
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="center")
        r0 = 2
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, Alignment(wrap_text=True, vertical="center", horizontal="center"), border
    ws.row_dimensions[r0].height = 32
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = bold if (first_bold and j == 1) else body
            c.alignment = center if isinstance(v, (int, float, dt.date)) and not isinstance(v, bool) else wrap
            c.border = border
            if (i - r0) % 2 == 0:
                c.fill = band
            if j in date_cols and v:
                c.number_format = "dd/mm/yyyy"
                c.alignment = center
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws, r0


def colorize(ws, col, r0, n, mapping):
    letter = get_column_letter(col)
    for k, fill in mapping.items():
        ws.conditional_formatting.add(f"{letter}{r0 + 1}:{letter}{r0 + n}",
                                      CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=fill)))


# ---------------------------------------------------------------- Léeme
ws = wb.active
ws.title = "Léeme"
ws.sheet_properties.tabColor = SALMON
lines = [
    ("Matriz Maestra de Reprogramación – Cronograma Integrado AFSP", Font(name=F, size=16, bold=True, color=AZUL)),
    ("Propuesta técnica de solución · Escenario 2 (Integración optimizada) · Corte del .mpp: 21/10/2026", Font(name=F, size=11, bold=True, color=SALMON)),
    ("", body),
    ("Cómo usar este libro", Font(name=F, size=12, bold=True, color=AZUL)),
    ("1. «Matriz maestra»: actividades que cambian, con fechas actuales y propuestas. Filtre por sistema, macroproceso, prioridad, responsable o estado.", body),
    ("2. «Dependencias»: predecesoras y sucesoras propuestas para Microsoft Project, con la holgura verificada.", body),
    ("3. «Importar a Project»: filas listas para cargar. También se entrega un archivo XML que Microsoft Project abre directamente.", body),
    ("4. «Validación diagnóstico», «Alternativas», «Escenarios», «Hitos», «Plan de acción», «Decisiones» y «Antes y después» sustentan la propuesta.", body),
    ("", body),
    ("Celdas editables", Font(name=F, size=12, bold=True, color=AZUL)),
    ("En «Matriz maestra», los CAF pueden ajustar las columnas en amarillo: Inicio propuesto, Fin propuesto, Prioridad, Responsable y Estado de la fecha. La variación se recalcula sola.", body),
    ("", body),
    ("Estados de la fecha", Font(name=F, size=12, bold=True, color=AZUL)),
    ("Calculada: obtenida con el calendario del .mpp y la red de vínculos.   Tentativa: requiere validar capacidad del CAF.   Pendiente de decisión: depende de una decisión de la Coordinación o la UCP.", body),
    ("", body),
    ("Advertencias", Font(name=F, size=12, bold=True, color=AZUL)),
    ("• Las fechas son de diseño funcional (MFD). La construcción del software (C3) no está programada en el .mpp, por lo que los hitos de software no son verificables.", body),
    ("• El archivo .mpp original no fue modificado. La hoja «Importar a Project» y el XML son la propuesta para que la PMO la cargue.", body),
    ("• Duraciones en días hábiles (calendario «Estándar» del .mpp: lunes a viernes, sin feriados registrados). Variaciones en días calendario.", body),
    ("• Las tareas nuevas usan códigos N# (tareas), M# (entregas parciales) y H# (hitos de integración).", body),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Matriz maestra
preds = {}
for l in links:
    lagtxt = "" if not l["desfase"] else (f"+{l['desfase']} días" if l["desfase"] > 0 else f"{l['desfase']} días")
    preds.setdefault(l["destino"], []).append(f"{l['origen']}{l['tipo']}{lagtxt}")
succs = {}
for l in links:
    succs.setdefault(l["origen"], []).append(l["destino"])

order = {"Alta": 0, "Media": 1, "Baja": 2}
rows = []
for it in sorted(items, key=lambda x: (x["id"].startswith("H"), order.get(x["prioridad"], 3), x["sistema"], x["ini_prop"])):
    rows.append([
        it["id"], it["sistema"], it["macro"], it["actividad"], it["tipo"],
        D(it["ini_act"]), D(it["fin_act"]), D(it["ini_prop"]), D(it["fin_prop"]), None, None,
        it["dur_prop"], "; ".join(preds.get(it["id"], [])), ", ".join(succs.get(it["id"], [])),
        it["caso"], it["prioridad"], it["responsable"], it["estado"], it["escenario"], it["justificacion"],
    ])
H = ["ID Project / código", "Sistema", "Macroproceso", "Actividad", "Tipo de cambio", "Inicio actual", "Fin actual",
     "Inicio propuesto", "Fin propuesto", "Variación inicio (días)", "Variación fin (días)", "Duración propuesta (días háb.)",
     "Predecesoras propuestas", "Sucesoras", "Caso", "Prioridad", "Responsable (CAF)", "Estado de la fecha", "Escenario", "Justificación"]
ws, r0 = sheet("Matriz maestra", H, rows, [11, 14, 15, 44, 20, 12, 12, 13, 13, 11, 11, 11, 26, 16, 9, 10, 26, 16, 9, 60],
               date_cols=(6, 7, 8, 9),
               note="Actividades que requieren modificación (Escenario 2). Columnas amarillas editables; la variación se recalcula. Filtros activos en el encabezado.")
n = len(rows)
yel = PatternFill("solid", fgColor="FFF4CC")
for i in range(r0 + 1, r0 + 1 + n):
    ws.cell(i, 10, f'=IF(AND(F{i}<>"",H{i}<>""),H{i}-F{i},"")').number_format = '+0;-0;0'
    ws.cell(i, 11, f'=IF(AND(G{i}<>"",I{i}<>""),I{i}-G{i},"")').number_format = '+0;-0;0'
    for j in (10, 11):
        ws.cell(i, j).font, ws.cell(i, j).alignment, ws.cell(i, j).border = Font(name=F, size=9, bold=True, color=SALMON), center, border
    for j in (8, 9, 16, 17, 18):
        ws.cell(i, j).fill = yel
dv1 = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"Calculada,Tentativa,Pendiente de decisión"', allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add(f"P{r0+1}:P{r0+n}"); dv2.add(f"R{r0+1}:R{r0+n}")
colorize(ws, 16, r0, n, LEVEL_FILL)

# resumen con fórmulas
ws2 = wb.create_sheet("Resumen", 1)
ws2.sheet_properties.tabColor = SALMON
ws2["A1"] = "Resumen de la reprogramación (se actualiza con la Matriz maestra)"
ws2["A1"].font = Font(name=F, size=14, bold=True, color=AZUL)
sistemas = ["Presupuesto", "Inversiones", "Abastecimiento", "GFRH", "Tesorería", "Endeudamiento", "Contabilidad", "Transversal"]
hdrs = ["Sistema", "Actividades", "Prioridad alta", "Calculadas", "Tentativas", "Pendientes de decisión", "Mayor variación de fin (días)"]
for j, h in enumerate(hdrs, 1):
    c = ws2.cell(3, j, h); c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, Alignment(wrap_text=True, horizontal="center", vertical="center"), border
MM = "'Matriz maestra'"
last = r0 + n
for i, s_ in enumerate(sistemas, 4):
    ws2.cell(i, 1, s_).font = bold
    ws2.cell(i, 2, f'=COUNTIF({MM}!$B${r0+1}:$B${last},A{i})')
    ws2.cell(i, 3, f'=COUNTIFS({MM}!$B${r0+1}:$B${last},A{i},{MM}!$P${r0+1}:$P${last},"Alta")')
    ws2.cell(i, 4, f'=COUNTIFS({MM}!$B${r0+1}:$B${last},A{i},{MM}!$R${r0+1}:$R${last},"Calculada")')
    ws2.cell(i, 5, f'=COUNTIFS({MM}!$B${r0+1}:$B${last},A{i},{MM}!$R${r0+1}:$R${last},"Tentativa")')
    ws2.cell(i, 6, f'=COUNTIFS({MM}!$B${r0+1}:$B${last},A{i},{MM}!$R${r0+1}:$R${last},"Pendiente de decisión")')
    ws2.cell(i, 7, f'=_xlfn.MAXIFS({MM}!$K${r0+1}:$K${last},{MM}!$B${r0+1}:$B${last},A{i})')
    for j in range(1, 8):
        ws2.cell(i, j).border = border
        if j > 1:
            ws2.cell(i, j).font, ws2.cell(i, j).alignment = body, center
tot = 4 + len(sistemas)
ws2.cell(tot, 1, "Total").font = Font(name=F, size=10, bold=True, color=SALMON)
for j in range(2, 7):
    L_ = get_column_letter(j)
    c = ws2.cell(tot, j, f"=SUM({L_}4:{L_}{tot-1})"); c.font, c.alignment, c.border = Font(name=F, size=10, bold=True, color=SALMON), center, border
ws2.cell(tot + 2, 1, "Nota: la columna «Mayor variación de fin» considera solo actividades existentes con fecha actual. "
                     "Las tareas nuevas no tienen variación.").font = Font(name=F, size=9, italic=True, color=GRIS)
for j, w in enumerate([18, 13, 13, 13, 13, 16, 18], 1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ---------------------------------------------------------------- Dependencias
rows = []
for l in links:
    rows.append([l["origen"], names.get(l["origen"], ""), l["destino"], names.get(l["destino"], ""), l["tipo"],
                 l["desfase"], l["holgura_dh"], "Sí" if l["cumple"] else "No", l["justificacion"], l["validacion"], l["ref"]])
ws, r0 = sheet("Dependencias", ["ID origen", "Actividad predecesora", "ID destino", "Actividad sucesora", "Tipo de relación",
                                "Desfase (días háb.)", "Holgura verificada (días háb.)", "Cumple con las fechas propuestas",
                                "Justificación", "Validación requerida", "Ref. diagnóstico"], rows,
               [9, 42, 9, 42, 9, 10, 12, 12, 48, 24, 12],
               note=f"{len(links)} vínculos propuestos. FC = Fin–Comienzo · CC = Comienzo–Comienzo · FF = Fin–Fin. Verificación: 0 conflictos y 0 ciclos con las fechas propuestas.")

# ---------------------------------------------------------------- Importar a Project
id_map = json.load(open(sys.argv[4])) if len(sys.argv) > 4 else {}
rows = []
for it in sorted(items, key=lambda x: x["ini_prop"]):
    rows.append([it["id"] if not it["id"][0].isalpha() else "", it["id"], id_map.get(it["id"], int(it["id"]) if it["id"].isdigit() else ""), it["actividad"], it.get("parent", ""),
                 D(it["ini_prop"]), D(it["fin_prop"]), f'{it["dur_prop"]:g} días' if it["dur_prop"] else "0 días",
                 "Sí" if it.get("milestone") else "No", "; ".join(preds.get(it["id"], [])),
                 "Manual (validar antes de pasar a automático)", it["justificacion"]])
sheet("Importar a Project", ["Id (existente)", "Código", "Id en el XML", "Nombre", "Tarea resumen sugerida (Id)", "Comienzo", "Fin", "Duración",
                             "Hito", "Predecesoras", "Modo de programación", "Notas"], rows,
      [9, 9, 9, 50, 12, 12, 12, 11, 7, 30, 22, 60], date_cols=(6, 7),
      note="Para la PMO: actualizar las filas con Id y crear las que solo tienen código bajo la tarea resumen indicada. "
           "Al crearlas, reemplazar los códigos N#, M# y H# de «Predecesoras» por los Id que asigne Project. "
           "Alternativa: abrir el archivo XML entregado, que ya trae todo cargado (columna «Id en el XML»).")

# ---------------------------------------------------------------- Validación diagnóstico
ws, r0 = sheet("Validación diagnóstico", ["Dependencia", "Descripción", "¿Obligatoria?", "Naturaleza", "Tratamiento", "Resultado", "Nota"],
               [list(x) for x in DEPS], [11, 50, 16, 20, 36, 12, 34],
               note="Revisión crítica de las 22 dependencias del diagnóstico y 2 nuevas (D23, D24).")
r = r0 + len(DEPS) + 3
ws.cell(r, 1, "Revisión de los 12 ajustes del diagnóstico").font = Font(name=F, size=12, bold=True, color=AZUL)
for j, h in enumerate(["Ajuste", "Ajuste original", "Decisión", "Motivo y reemplazo"], 1):
    c = ws.cell(r + 1, j, h); c.font, c.fill, c.border, c.alignment = hdr_font, hdr_fill, border, Alignment(wrap_text=True)
for i, a in enumerate(AJUSTES, r + 2):
    for j, v in enumerate(a, 1):
        c = ws.cell(i, j, v); c.font, c.border, c.alignment = (bold if j == 1 else body), border, wrap
r = r + 3 + len(AJUSTES)
ws.cell(r, 1, "Correcciones al diagnóstico").font = Font(name=F, size=12, bold=True, color=AZUL)
for j, h in enumerate(["Tema", "Qué se corrige", "Resultado"], 1):
    c = ws.cell(r + 1, j, h); c.font, c.fill, c.border, c.alignment = hdr_font, hdr_fill, border, Alignment(wrap_text=True)
for i, a in enumerate(CORRECCIONES, r + 2):
    for j, v in enumerate(a, 1):
        c = ws.cell(i, j, v); c.font, c.border, c.alignment = (bold if j == 1 else body), border, wrap

# ---------------------------------------------------------------- Alternativas
rows = []
for caso, alts, rec, why in ALTERNATIVAS:
    for code, desc, ev in alts:
        rows.append([caso, code, desc] + ev + ["Recomendada" if code == rec else "", why if code == rec else ""])
ws, r0 = sheet("Alternativas", ["Caso", "Alt.", "Descripción"] + CRIT + ["Selección", "Por qué"], rows,
               [26, 6, 44] + [13] * 8 + [13, 46], note="Evaluación de alternativas A, B y C por caso. Se prioriza la integridad funcional, no el menor movimiento de fechas.")
colorize(ws, 12, r0, len(rows), {"Recomendada": "D8F0EE"})

# ---------------------------------------------------------------- Escenarios
ws, r0 = sheet("Escenarios", ["Criterio", "Escenario 1 – Conservador", "Escenario 2 – Integración optimizada (recomendado)", "Escenario 3 – Integración prioritaria"],
               [list(x) for x in ESCENARIOS], [30, 40, 46, 40],
               note="Comparación de escenarios. «No verificable» indica que el .mpp no tiene datos (recursos, construcción) para calcularlo.")
for i in range(r0 + 1, r0 + 1 + len(ESCENARIOS)):
    ws.cell(i, 3).fill = PatternFill("solid", fgColor="D8F0EE")

# ---------------------------------------------------------------- Hitos
hd = plan["hitos"]
rows = [[h[0], h[1], h[2], h[3], D(hd[h[0]]), h[4], h[5], "No verificable: el .mpp no programa la construcción (C3)"] for h in HITOS]
sheet("Hitos de integración", ["Hito", "Descripción", "Sistemas involucrados", "Entregables necesarios", "Fecha objetivo (diseño)",
                               "Criterio de aceptación", "Responsable propuesto", "Hito de software y puesta en uso"], rows,
      [6, 38, 30, 32, 13, 46, 24, 30], date_cols=(5,),
      note="Hitos de aprobación del diseño funcional integrado. No se declara alcanzado un hito sin el acta o evidencia del criterio de aceptación.")

# ---------------------------------------------------------------- Evaluación transversal
sheet("Evaluación transversal", ["Resultado a evaluar", "Sistema que produce los datos", "Proceso responsable de evaluarlo"],
      [list(x) for x in EVAL_TRANS], [40, 30, 70], note="Caso 7: no se crea un módulo de Evaluación por sistema.")

# ---------------------------------------------------------------- Plan de acción
ws, r0 = sheet("Plan de acción", ["Prioridad", "Tipo de acción", "Acción", "Sistema", "CAF o responsable por validar", "Fecha límite propuesta",
                                  "Entregable", "Nivel de decisión"], [list(x) for x in PLAN], [9, 18, 56, 26, 32, 14, 24, 14],
               note="Plan de implementación de la solución. Nivel de decisión: CAF, Varios CAF, Coordinación o UCP.")

# ---------------------------------------------------------------- Decisiones
sheet("Decisiones", ["Código", "Decisión", "Problema que resuelve", "Alternativas disponibles", "Alternativa recomendada",
                     "Consecuencias de no decidir", "Sistemas afectados", "Responsable de aprobación", "Fecha límite propuesta"],
      [list(x) for x in DECISIONES], [9, 28, 38, 52, 12, 40, 24, 26, 13],
      note="Entregable 4 · Matriz de decisiones para la Coordinación del Proyecto y la UCP.")

# ---------------------------------------------------------------- Antes y después
ws, r0 = sheet("Antes y después", ["Hallazgo original", "Riesgo inicial", "Solución propuesta", "Nueva fecha", "Riesgo residual", "Estado de solución"],
               [list(x) for x in ANTES_DESPUES], [48, 12, 42, 30, 12, 22],
               note="Matriz de validación antes y después. No se declara resuelto lo que depende de una decisión o de recursos no confirmados.")
colorize(ws, 2, r0, len(ANTES_DESPUES), LEVEL_FILL)
colorize(ws, 5, r0, len(ANTES_DESPUES), LEVEL_FILL)
colorize(ws, 6, r0, len(ANTES_DESPUES), STATE_FILL)

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save(OUT)
print("ok")
